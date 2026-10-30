"""Regression checks for fixture identity, source joins and partial imports."""
import io
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from flask import Flask
from server_core.routes.fitness_data_api import fitness_data_bp, _build_player_doc_from_team_overview
from server_core.services.fitness_history import extract_match


def fixture(distance=1000, valid_date=True):
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A3"], cover["A5"] = "甲 - 乙", "中超"
    cover["A7"] = datetime(2026, 7, 26) if valid_date else ""
    for side, team in (("Home", "甲"), ("Away", "乙")):
        ws = wb.create_sheet(f"{side} Team Overview")
        ws["B3"] = team
        ws.append([None])
        ws.append([None, "球员", "出场时间", "总距离 (米)", "平均速度 (米/秒)"])
        for i, name in enumerate(("泽卡", "球员二", "球员三"), 1):
            ws.append([None, f"{i} - {name}", 90, distance, 2])
        acc = wb.create_sheet(f"{side} High Accel Decel")
        acc.append([None, "球员", "出场时间", "总计", "高加速", "高减速"])
        acc.append([None, "1 - 泽卡", 90, 12, 0, 12])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


class FitnessHistoryTests(unittest.TestCase):
    def test_sources_and_identity(self):
        a = extract_match(fixture(), "a.xlsx", _build_player_doc_from_team_overview)
        b = extract_match(fixture(2000), "renamed.xlsx", _build_player_doc_from_team_overview)
        self.assertEqual(a["id"], b["id"])
        zeca = [p for p in a["players"] if p["player"] == "泽卡"]
        self.assertEqual(len(zeca), 2)
        self.assertEqual({p["team"] for p in zeca}, {"甲", "乙"})
        self.assertEqual(zeca[0]["raw"]["高加速次数 (3–10 米/秒²)"], 0)
        self.assertEqual(zeca[0]["raw"]["高减速次数 (-3至-10 米/秒²)"], 12)

    def test_missing_date_rejected(self):
        with self.assertRaisesRegex(ValueError, "日期"):
            extract_match(fixture(valid_date=False), "a.xlsx", _build_player_doc_from_team_overview)

    def test_partial_import_and_replacement(self):
        app = Flask(__name__)
        app.register_blueprint(fitness_data_bp)
        with tempfile.TemporaryDirectory() as tmp, patch("server_core.routes.fitness_data_api.user_data_file", side_effect=lambda user, name: Path(tmp) / name):
            client = app.test_client()
            result = client.post("/api/fitness-data/history", data={"files": [(fixture(), "a.xlsx"), (io.BytesIO(b"bad"), "bad.xlsx")]}).get_json()
            self.assertEqual(result["imported"], 1)
            self.assertEqual(len(result["errors"]), 1)
            result = client.post("/api/fitness-data/history", data={"files": (fixture(2000), "new.xlsx")}).get_json()
            self.assertEqual(len(result["matches"]), 1)
            self.assertEqual(result["matches"][0]["players"][0]["raw"]["总距离 (米)"], 2000)
            self.assertTrue((Path(tmp) / "fitness_history.json.bak").exists())
            self.assertEqual(len(client.get("/api/fitness-data/history").get_json()["matches"]), 1)


if __name__ == "__main__":
    unittest.main()
