"""Extract match-level player records without changing the single-match importer."""
import hashlib
import re
from datetime import datetime, date
from openpyxl import load_workbook


def extract_match(stream, filename, parse_players):
    content = stream.read()
    from io import BytesIO
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        if "Cover" not in wb.sheetnames:
            raise ValueError("缺少 Cover 比赛信息表")
        cover = list(wb["Cover"].iter_rows(values_only=True))
        def cell(row):
            return cover[row - 1][0] if len(cover) >= row else None
        match_name = str(cell(3) or filename)
        competition = str(cell(5) or "未注明赛事")
        played = cell(7)
        if isinstance(played, (datetime, date)):
            match_date = played.isoformat()[:10]
        else:
            try:
                match_date = date.fromisoformat(str(played or "").strip()[:10]).isoformat()
            except ValueError:
                raise ValueError("Cover 第7行缺少有效比赛日期（需要Excel日期或YYYY-MM-DD）")
        teams = re.split(r"\s+-\s+", match_name, maxsplit=1)
        records = []
        for side, team_index in (("Home", 0), ("Away", 1)):
            sheet = f"{side} Team Overview"
            if sheet not in wb.sheetnames:
                continue
            rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
            doc, error = parse_players(rows)
            if doc is None:
                raise ValueError(f"{sheet}: {error or '未识别到球员'}")
            team = str(rows[2][1] or (teams[team_index] if len(teams) > team_index else side))
            opponent = teams[1 - team_index] if len(teams) == 2 else ""
            acceleration = {}
            accel_sheet = f"{side} High Accel Decel"
            if accel_sheet in wb.sheetnames:
                accel_rows = list(wb[accel_sheet].iter_rows(values_only=True))
                headers = next((r for r in accel_rows if "高加速" in r and "高减速" in r), None)
                accel_index = headers.index("高加速") if headers else -1
                decel_index = headers.index("高减速") if headers else -1
                for row in accel_rows if headers else []:
                    if len(row) > 5 and re.match(r"^\d+\s*-\s*", str(row[1] or "")):
                        acceleration[str(row[1]).strip()] = {"高加速次数 (3–10 米/秒²)": row[accel_index], "高减速次数 (-3至-10 米/秒²)": row[decel_index]}
            for player in doc["players"]:
                name = re.sub(r"^\s*\d+\s*[-–.]\s*", "", player["player"]).strip()
                if name.lower() in ("total", "average", "合计", "平均", "平均值", "总计"):
                    continue
                records.append({"player": name, "sourcePlayer": player["player"], "team": team,
                                "opponent": opponent, "side": side.lower(),
                                "raw": {**player["raw"], **acceleration.get(player["player"].strip(), {})}})
        if not records:
            raise ValueError("未找到 Home/Away Team Overview 球员数据")
        # A new version of the same fixture replaces its previous records.
        key = hashlib.sha256(f"{competition}|{match_date}|{match_name}".encode()).hexdigest()[:20]
        return {"id": key, "match": match_name, "date": match_date, "competition": competition,
                "sourceFile": filename, "players": records}
    finally:
        wb.close()
